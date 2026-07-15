"""
Shared DOCX style definitions matching reference report format.

Reference: 平远县泗水镇金田村村庄规划（2022-2035年）文本.docx

Format spec:
  一级标题: 仿宋/TNR 18pt Bold Center, single line spacing
  二级标题: 仿宋/TNR 14pt NOT bold Left, single line spacing
  三级标题: 仿宋/TNR 12pt Bold JUSTIFY, 1.5x line spacing
  正文:     仿宋/TNR 12pt JUSTIFY, 1.5x line spacing, ~2 char indent
  表格表头: 仿宋 ~11pt Bold Center
  表格数据: 仿宋 10.5pt Center
  表题/图题: 宋体/TNR 10.5pt Center
  页边距: 左3.17cm 右3.17cm 上2.54cm 下2.54cm
"""

from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Font constants
CN_FONT = '仿宋'
EN_FONT = 'Times New Roman'
CN_FONT_TABLE_TITLE = '宋体'

# Page margins (reference report)
PAGE_MARGIN_LEFT = Cm(3.17)
PAGE_MARGIN_RIGHT = Cm(3.17)
PAGE_MARGIN_TOP = Cm(2.54)
PAGE_MARGIN_BOTTOM = Cm(2.54)


def _set_style_east_asia(style, cn_font, en_font=EN_FONT):
    """Set eastAsia and Western fonts on a style definition."""
    rpr = style.element.find(qn('w:rPr'))
    if rpr is None:
        rpr = OxmlElement('w:rPr')
        style.element.append(rpr)
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:eastAsia'), cn_font)
    rfonts.set(qn('w:ascii'), en_font)
    rfonts.set(qn('w:hAnsi'), en_font)


def set_run_font(run, cn_font=CN_FONT, en_font=EN_FONT, font_size=12, bold=False):
    """Set run font with both CJK and Western font names."""
    run.font.name = en_font
    run.font.size = Pt(font_size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:eastAsia'), cn_font)
    rfonts.set(qn('w:ascii'), en_font)
    rfonts.set(qn('w:hAnsi'), en_font)


def set_cell_font(cell, cn_font=CN_FONT, en_font=EN_FONT, font_size=10.5, bold=False):
    """Set cell font with both CJK and Western font names, centered, single line spacing."""
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        for run in paragraph.runs:
            set_run_font(run, cn_font, en_font, font_size, bold)


def setup_document_styles(doc):
    """Configure all built-in styles to match reference report format.

    Must be called after Document() creation, before adding content.
    """
    # Title (level=0): 仿宋/TNR 24pt Bold Center
    title_style = doc.styles['Title']
    title_style.font.size = Pt(24)
    title_style.font.bold = True
    title_style.font.name = EN_FONT
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.first_line_indent = Pt(0)
    _set_style_east_asia(title_style, CN_FONT)

    # Heading 1 (章标题): 仿宋/TNR 18pt Bold Center, single line spacing
    h1 = doc.styles['Heading 1']
    h1.font.name = EN_FONT
    h1.font.size = Pt(18)
    h1.font.bold = True
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.first_line_indent = Pt(0)
    h1.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    h1.paragraph_format.space_before = Pt(24)
    h1.paragraph_format.space_after = Pt(24)
    _set_style_east_asia(h1, CN_FONT)

    # Heading 2 (条标题): 仿宋/TNR 14pt NOT bold Left, single line spacing
    h2 = doc.styles['Heading 2']
    h2.font.name = EN_FONT
    h2.font.size = Pt(14)
    h2.font.bold = False
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.first_line_indent = Pt(0)
    h2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(12)
    _set_style_east_asia(h2, CN_FONT)

    # Heading 3 (小标题): 仿宋/TNR 12pt Bold JUSTIFY, 1.5x line spacing
    h3 = doc.styles['Heading 3']
    h3.font.name = EN_FONT
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    h3.paragraph_format.first_line_indent = Pt(24)
    h3.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    _set_style_east_asia(h3, CN_FONT)

    # Normal (正文): 仿宋/TNR 12pt JUSTIFY, 1.5x line spacing
    normal = doc.styles['Normal']
    normal.font.name = EN_FONT
    normal.font.size = Pt(12)
    normal.font.bold = False
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_before = Pt(12)
    normal.paragraph_format.space_after = Pt(0)
    _set_style_east_asia(normal, CN_FONT)


def setup_page_margins(doc):
    """Set page margins matching reference report."""
    for section in doc.sections:
        section.left_margin = PAGE_MARGIN_LEFT
        section.right_margin = PAGE_MARGIN_RIGHT
        section.top_margin = PAGE_MARGIN_TOP
        section.bottom_margin = PAGE_MARGIN_BOTTOM


def fix_heading_fonts(heading, cn_font=CN_FONT, en_font=EN_FONT):
    """Ensure eastAsia font is set on heading runs.

    python-docx add_heading() may create runs without rPr or rFonts,
    causing CJK characters to render in SimSun instead of FangSong.
    """
    for run in heading.runs:
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is not None:
            rfonts.set(qn('w:eastAsia'), cn_font)
            rfonts.set(qn('w:hAnsi'), en_font)
            if not rfonts.get(qn('w:ascii')):
                rfonts.set(qn('w:ascii'), en_font)
        else:
            rfonts = OxmlElement('w:rFonts')
            rfonts.set(qn('w:eastAsia'), cn_font)
            rfonts.set(qn('w:ascii'), en_font)
            rfonts.set(qn('w:hAnsi'), en_font)
            rpr.insert(0, rfonts)
