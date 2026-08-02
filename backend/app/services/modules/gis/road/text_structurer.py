"""Structure road-planning evidence from workbooks and planning text.

The extractor is deliberately conservative: it uses deterministic parsing,
keeps every source locator, performs exact-only place matching, and emits a
blocking review state whenever a road endpoint cannot be located uniquely.
Planned road/project geometries are never accepted as location sources.
"""

from __future__ import annotations

import hashlib
import json
import re
from collections import defaultdict
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, MutableMapping, Optional, Sequence, Tuple

from .schemas import Stage2StructureConfig, Stage2StructureResult


SCHEMA_VERSION = "1.0"
NUMBER_MARKERS = "①②③④⑤⑥⑦⑧⑨⑩"
ROAD_TERMS = re.compile(r"道路|村道|巷道|机耕路|步道|绿道|公路|乡道")
DEFICIENCY_TERMS = re.compile(r"未硬化|荒废|没有|无专门|不足|散乱|缺少|亟需|急需")
STRICT_TERMS = re.compile(r"严禁|禁止|不得|不准|不占用|不能|严格限制|禁止开发")


def _clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").replace("_", " ")
    return re.sub(r"\s+", " ", text).strip()


def _norm(value: Any) -> str:
    return re.sub(r"[\s_（）()、,，;；:：\-—–/]+", "", _clean(value)).lower()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _stable_id(prefix: str, *parts: Any) -> str:
    payload = "\x1f".join(_clean(part) for part in parts)
    return prefix + hashlib.sha1(payload.encode("utf-8")).hexdigest()[:12].upper()


def _write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def _relative(path: Path) -> str:
    try:
        return path.resolve().relative_to(Path.cwd().resolve()).as_posix()
    except ValueError:
        return path.resolve().as_posix()


def _evidence(
    path: Path,
    source_text: str,
    *,
    source_kind: str,
    confidence: float,
    sheet: Optional[str] = None,
    row: Optional[int] = None,
    cell_range: Optional[str] = None,
    line: Optional[int] = None,
    ambiguous: bool = False,
    ambiguity_reason: Optional[str] = None,
    requires_confirmation: bool = False,
) -> Dict[str, Any]:
    document = _relative(path)
    locator = cell_range or (f"line:{line}" if line is not None else "document")
    evidence_id = _stable_id("EV", document, locator, source_text)
    return {
        "evidence_id": evidence_id,
        "source_document": document,
        "source_kind": source_kind,
        "source_sheet": sheet,
        "source_row": row,
        "source_cell_range": cell_range,
        "source_line": line,
        "source_text": _clean(source_text),
        "extraction_method": "deterministic_rule_v1",
        "extraction_confidence": round(float(confidence), 3),
        "is_ambiguous": bool(ambiguous),
        "ambiguity_reason": ambiguity_reason,
        "requires_human_confirmation": bool(requires_confirmation),
    }


def _row_range(sheet: Any, row_number: int) -> str:
    cells = [cell for cell in sheet[row_number] if _clean(cell.value)]
    if not cells:
        return f"{sheet.title}!{row_number}:{row_number}"
    return f"{sheet.title}!{cells[0].coordinate}:{cells[-1].coordinate}"


def _load_workbook(path: Path) -> Any:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - environment-specific error path
        raise RuntimeError("Stage 2 requires openpyxl>=3.1 for read-only XLSX extraction") from exc
    return openpyxl.load_workbook(path, read_only=False, data_only=True)


def _find_header(nonempty: Sequence[Tuple[int, str]], keyword: str) -> Optional[int]:
    for column, value in nonempty:
        if keyword in value:
            return column
    return None


def _is_header_row(nonempty: Sequence[Tuple[int, str]]) -> bool:
    values = "|".join(value for _, value in nonempty)
    return "设施名称" in values and "所在空间位置" in values and "备注" in values


def _looks_like_metric(value: str) -> bool:
    compact = _clean(value)
    return bool(
        re.fullmatch(r"[\d.\s~～多余约]+(?:m|km|米|千米|㎡|㎞|亩|公顷|个|处|段)?", compact, re.I)
        or compact in {"①", "②", "——", "-"}
    )


def _classify_demand(name: str, category: str = "") -> Optional[str]:
    text = f"{category} {name}"
    if re.search(r"幼儿园|小学|学校|教育|研学", text):
        return "education"
    if re.search(r"卫生站|卫生室|医疗|诊所", text):
        return "healthcare"
    if re.search(r"养老|敬老", text):
        return "eldercare"
    if re.search(r"防震|避灾|疏散|消防|应急", text):
        return "emergency"
    if re.search(r"民宿|观景|旅游|体验园|公园|营地|采摘", text):
        return "tourism"
    if re.search(r"产业|加工|合作社|基地|农场|种植园|集市|养殖|药苗", text):
        return "industry"
    if re.search(r"公交站|停车场", text):
        return "transport_facility"
    if re.search(r"服务中心|文化|文体|篮球|健身|活动场地|菜市场|广场", text):
        return "public_service"
    if re.search(r"污水|垃圾|公厕|供水|供电|变压器", text):
        return "utility"
    if re.search(r"祠堂|古驿道|古茶亭|历史", text):
        return "cultural"
    return None


def _split_places(value: str) -> List[str]:
    text = _clean(value)
    text = re.sub(r"^(?:经过|位于)", "", text)
    if not text:
        return []
    parts = [_clean(part) for part in re.split(r"[、,，;；]", text)]
    return [part for part in parts if part and part not in {"——", "无"}]


def _source_row_text(values: Sequence[Any]) -> str:
    return "；".join(_clean(value) for value in values if _clean(value))


def _extract_workbook_rows(path: Path, role: str) -> List[Dict[str, Any]]:
    workbook = _load_workbook(path)
    extracted: List[Dict[str, Any]] = []
    for sheet in workbook.worksheets:
        header: Optional[Dict[str, int]] = None
        current_category = ""
        for row_number in range(1, sheet.max_row + 1):
            values = [cell.value for cell in sheet[row_number]]
            nonempty = [(index + 1, _clean(value)) for index, value in enumerate(values) if _clean(value)]
            if _is_header_row(nonempty):
                header = {
                    "type": _find_header(nonempty, "类型") or 2,
                    "facility": _find_header(nonempty, "设施名称") or 3,
                    "independent": _find_header(nonempty, "是否独立用地") or sheet.max_column,
                    "location": _find_header(nonempty, "所在空间位置") or sheet.max_column,
                    "service": _find_header(nonempty, "服务对象") or sheet.max_column,
                    "coverage": _find_header(nonempty, "覆盖程度") or sheet.max_column,
                    "note": _find_header(nonempty, "备注") or sheet.max_column,
                }
                current_category = ""
                continue
            if header is None or not nonempty:
                continue

            category_value = _clean(values[header["type"] - 1])
            if category_value:
                current_category = category_value
            name_values = []
            for column in range(header["facility"], header["independent"]):
                text = _clean(values[column - 1])
                if text and not _looks_like_metric(text):
                    name_values.append(text)
            name = " / ".join(dict.fromkeys(name_values))
            if not name or name.startswith("注：") or _norm(name) in {"其他", "传统古村落是否"}:
                continue

            location = _clean(values[header["location"] - 1])
            note = _clean(values[header["note"] - 1])
            if role == "current" and not DEFICIENCY_TERMS.search(note):
                continue
            if not location and note in {"", "无", "——"}:
                continue
            if "××" in name and not location:
                continue
            extracted.append(
                {
                    "source_path": path,
                    "source_sheet": sheet.title,
                    "source_row": row_number,
                    "source_cell_range": _row_range(sheet, row_number),
                    "source_text": _source_row_text(values),
                    "role": role,
                    "category": current_category,
                    "name": name,
                    "location": location,
                    "service_object": _clean(values[header["service"] - 1]),
                    "coverage": _clean(values[header["coverage"] - 1]),
                    "note": note,
                }
            )
    workbook.close()
    return extracted


class _TableHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.rows: List[List[Tuple[str, int, int]]] = []
        self._row: Optional[List[Tuple[str, int, int]]] = None
        self._cell: Optional[List[str]] = None
        self._rowspan = 1
        self._colspan = 1

    def handle_starttag(self, tag: str, attrs: List[Tuple[str, Optional[str]]]) -> None:
        if tag == "tr":
            self._row = []
        elif tag in {"td", "th"} and self._row is not None:
            attributes = dict(attrs)
            self._rowspan = int(attributes.get("rowspan") or 1)
            self._colspan = int(attributes.get("colspan") or 1)
            self._cell = []

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag in {"td", "th"} and self._cell is not None and self._row is not None:
            self._row.append((_clean("".join(self._cell)), self._rowspan, self._colspan))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            self.rows.append(self._row)
            self._row = None


def _expand_html_rows(rows: Sequence[Sequence[Tuple[str, int, int]]]) -> List[List[str]]:
    pending: Dict[int, Tuple[str, int]] = {}
    output: List[List[str]] = []
    for raw_row in rows:
        grid: Dict[int, str] = {}
        for column, (value, remaining) in list(pending.items()):
            grid[column] = value
            if remaining <= 1:
                del pending[column]
            else:
                pending[column] = (value, remaining - 1)
        column = 0
        for value, rowspan, colspan in raw_row:
            while column in grid:
                column += 1
            for offset in range(colspan):
                grid[column + offset] = value
                if rowspan > 1:
                    pending[column + offset] = (value, rowspan - 1)
            column += colspan
        width = max(grid, default=-1) + 1
        output.append([grid.get(index, "") for index in range(width)])
    max_width = max((len(row) for row in output), default=0)
    return [row + [""] * (max_width - len(row)) for row in output]


def _markdown_project_rows(paths: Sequence[Path]) -> List[Dict[str, Any]]:
    results: List[Dict[str, Any]] = []
    for path in paths:
        for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
            if "<table" not in line or "项目名称" not in line or "空间位置" not in line:
                continue
            parser = _TableHTMLParser()
            parser.feed(line)
            grid = _expand_html_rows(parser.rows)
            header_index = next(
                (index for index, row in enumerate(grid) if any("项目名称" in cell for cell in row)),
                None,
            )
            if header_index is None:
                continue
            headers = grid[header_index]
            name_col = next(index for index, value in enumerate(headers) if "项目名称" in value)
            location_col = next(index for index, value in enumerate(headers) if "空间位置" in value)
            category_col = next(
                (index for index, value in enumerate(headers) if "项目类型" in value or "项目分类" in value),
                0,
            )
            sequence_col = next((index for index, value in enumerate(headers) if "序号" in value), None)
            timeline_col = next((index for index, value in enumerate(headers) if "时序" in value), None)
            for row in grid[header_index + 1 :]:
                name = _clean(row[name_col] if name_col < len(row) else "")
                location = _clean(row[location_col] if location_col < len(row) else "")
                if not name or name in {"——", "无"}:
                    continue
                results.append(
                    {
                        "source_path": path,
                        "source_line": line_number,
                        "source_text": "；".join(value for value in row if value),
                        "category": _clean(row[category_col] if category_col < len(row) else ""),
                        "name": name,
                        "location": location,
                        "sequence": _clean(row[sequence_col]) if sequence_col is not None else "",
                        "timeline": _clean(row[timeline_col]) if timeline_col is not None else "",
                    }
                )
    return results


def parse_numbered_segments(value: str) -> List[Tuple[str, str]]:
    """Return ``(marker, segment)`` pairs without inventing missing markers."""

    text = _clean(value)
    matches = list(re.finditer(f"([{NUMBER_MARKERS}])", text))
    if not matches:
        return [("", text)] if text else []
    output: List[Tuple[str, str]] = []
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        output.append((match.group(1), _clean(text[match.end() : end]).strip("；;")))
    return output


def _route_parts(value: str) -> Tuple[Optional[str], Optional[str], List[str], bool]:
    text = re.sub(r"^\d+", "", _clean(value)).strip()
    pieces = _split_places(text)
    route_piece = next((piece for piece in pieces if re.search(r"--?|—|–|至|到", piece)), None)
    if route_piece:
        endpoints = [_clean(item) for item in re.split(r"--?|—|–|至|到", route_piece, maxsplit=1)]
        if len(endpoints) == 2 and all(endpoints):
            return endpoints[0], endpoints[1], pieces, False
    return None, None, pieces, True


def _action(value: str) -> str:
    if "拓宽" in value:
        return "widen"
    if "新增" in value or "新建" in value or "建设" in value:
        return "new"
    if "硬化" in value or "硬底化" in value:
        return "harden"
    if "修复" in value or "修缮" in value:
        return "repair"
    if "提升" in value or "改善" in value:
        return "improve"
    return "unspecified"


def _road_class(name: str) -> str:
    name = _norm(name)
    if "巷道" in name:
        return "lane"
    if "机耕" in name or "农田步道" in name:
        return "farm_access"
    if "滨水" in name:
        return "waterfront_trail"
    if "森林" in name:
        return "forest_trail"
    if "绿道" in name:
        return "greenway"
    if "乡道" in name:
        return "township_road"
    if "村道" in name:
        return "village_road"
    return "road_unspecified"


def _dimensions(text: str) -> Dict[str, Any]:
    text = re.sub(r"\s+", "", text)
    result: Dict[str, Any] = {
        "min_width_m": None,
        "existing_width_m": None,
        "proposed_length_m": None,
        "proposed_length_min_m": None,
        "proposed_length_max_m": None,
    }
    widen = re.search(r"由\s*(\d+(?:\.\d+)?)\s*m?\s*拓宽到\s*(\d+(?:\.\d+)?)\s*m", text, re.I)
    if widen:
        result["existing_width_m"] = float(widen.group(1))
        result["min_width_m"] = float(widen.group(2))
    else:
        width = re.search(r"(?:新增|新建|硬化|建设)?\s*(\d+(?:\.\d+)?)\s*m\s*宽", text, re.I)
        if width:
            result["min_width_m"] = float(width.group(1))
    length_range = re.search(r"(\d+(?:\.\d+)?)\s*[~～-]\s*(\d+(?:\.\d+)?)\s*(km|千米|m|米)", text, re.I)
    if length_range:
        factor = 1000.0 if length_range.group(3).lower() in {"km", "千米"} else 1.0
        result["proposed_length_min_m"] = float(length_range.group(1)) * factor
        result["proposed_length_max_m"] = float(length_range.group(2)) * factor
    else:
        length = re.search(r"(?:长|长度[:：]?)(?:约)?\s*(\d+(?:\.\d+)?)\s*(km|千米|m|米)", text, re.I)
        if not length:
            length = re.search(r"(\d+(?:\.\d+)?)\s*(km|千米|m|米)\s*长", text, re.I)
        if length:
            factor = 1000.0 if length.group(2).lower() in {"km", "千米"} else 1.0
            result["proposed_length_m"] = float(length.group(1)) * factor
    return result


class ExactLocationIndex:
    """Exact-only spatial lookup; fuzzy coordinate guessing is prohibited."""

    NAME_FIELDS = ("name", "NAME", "名称", "注记", "label", "text")

    def __init__(self, sources: Sequence[Path]) -> None:
        self.sources = tuple(sources)
        self._items: MutableMapping[str, List[Dict[str, Any]]] = defaultdict(list)
        for path in sources:
            payload = json.loads(path.read_text(encoding="utf-8-sig"))
            for feature_index, feature in enumerate(payload.get("features", [])):
                geometry = feature.get("geometry")
                if not geometry or geometry.get("type") != "Point":
                    continue
                properties = feature.get("properties") or {}
                for field in self.NAME_FIELDS:
                    name = _clean(properties.get(field))
                    if len(_norm(name)) < 2:
                        continue
                    self._items[_norm(name)].append(
                        {
                            "name": name,
                            "geometry": geometry,
                            "source": _relative(path),
                            "feature_index": feature_index,
                            "matched_field": field,
                        }
                    )

    def resolve(self, place_name: str) -> Dict[str, Any]:
        matches = self._items.get(_norm(place_name), []) if place_name else []
        unique: Dict[str, Dict[str, Any]] = {}
        for match in matches:
            key = json.dumps(match["geometry"], sort_keys=True, ensure_ascii=False)
            unique.setdefault(key, match)
        matches = list(unique.values())
        if len(matches) == 1:
            return {
                "location_status": "matched",
                "location_confidence": 1.0,
                "geometry": matches[0]["geometry"],
                "location_match": matches[0],
                "location_reason": "authoritative source contained one exact name match",
            }
        if len(matches) > 1:
            return {
                "location_status": "ambiguous",
                "location_confidence": 0.0,
                "geometry": None,
                "location_match": None,
                "location_candidates": matches,
                "location_reason": "more than one exact spatial match was found",
            }
        return {
            "location_status": "unresolved",
            "location_confidence": 0.0,
            "geometry": None,
            "location_match": None,
            "location_reason": "地图和允许使用的台账中未找到唯一对应位置",
        }


def _make_demand(
    *,
    name: str,
    demand_type: str,
    place_name: str,
    priority: str,
    evidence: Dict[str, Any],
    location_index: ExactLocationIndex,
    service_object: str = "",
    coverage: str = "",
    source_role: str = "planned_demand",
) -> Dict[str, Any]:
    location = location_index.resolve(place_name or name)
    requires_confirmation = location["location_status"] != "matched"
    evidence = dict(evidence)
    evidence["requires_human_confirmation"] = (
        evidence["requires_human_confirmation"] or requires_confirmation
    )
    return {
        "name": name,
        "demand_type": demand_type,
        "place_name": place_name or None,
        "population_served": None,
        "service_object": service_object or None,
        "coverage": coverage or None,
        "priority": priority,
        "source_role": source_role,
        **location,
        "is_ambiguous": location["location_status"] == "ambiguous",
        "requires_human_confirmation": requires_confirmation,
        "evidence": evidence,
        "supporting_evidence": [],
    }


def _workbook_demands(
    rows: Sequence[Mapping[str, Any]], location_index: ExactLocationIndex
) -> List[Dict[str, Any]]:
    demands: List[Dict[str, Any]] = []
    for row in rows:
        demand_type = _classify_demand(str(row["name"]), str(row["category"]))
        if demand_type is None or ROAD_TERMS.search(_norm(row["name"])):
            continue
        evidence = _evidence(
            Path(row["source_path"]),
            str(row["source_text"]),
            source_kind="xlsx_row",
            confidence=0.98 if row["role"] == "demand" else 0.9,
            sheet=str(row["source_sheet"]),
            row=int(row["source_row"]),
            cell_range=str(row["source_cell_range"]),
        )
        places = _split_places(str(row["location"])) or [""]
        for place in places:
            display_name = str(row["name"])
            if len(places) > 1:
                display_name = f"{display_name}（{place}）"
            demands.append(
                _make_demand(
                    name=display_name,
                    demand_type=demand_type,
                    place_name=place,
                    priority="high" if row["role"] == "demand" else "medium",
                    evidence=evidence,
                    location_index=location_index,
                    service_object=str(row["service_object"]),
                    coverage=str(row["coverage"]),
                    source_role="planned_demand" if row["role"] == "demand" else "current_deficiency",
                )
            )
    return demands


def _markdown_demands(
    rows: Sequence[Mapping[str, Any]], location_index: ExactLocationIndex
) -> List[Dict[str, Any]]:
    demands: List[Dict[str, Any]] = []
    for row in rows:
        demand_type = _classify_demand(str(row["name"]), str(row["category"]))
        if demand_type is None or ROAD_TERMS.search(_norm(row["name"])):
            continue
        evidence = _evidence(
            Path(row["source_path"]),
            str(row["source_text"]),
            source_kind="markdown_table_row",
            confidence=0.95,
            line=int(row["source_line"]),
        )
        places = _split_places(str(row["location"])) or [""]
        for place in places:
            display_name = str(row["name"])
            if len(places) > 1:
                display_name = f"{display_name}（{place}）"
            demands.append(
                _make_demand(
                    name=display_name,
                    demand_type=demand_type,
                    place_name=place,
                    priority="high" if "2022-2025" in str(row.get("timeline", "")) else "medium",
                    evidence=evidence,
                    location_index=location_index,
                )
            )
    return demands


def _project_from_parts(
    *,
    name: str,
    location_text: str,
    requirement_text: str,
    priority: str,
    evidence: Dict[str, Any],
    location_index: ExactLocationIndex,
    marker: str = "",
) -> Dict[str, Any]:
    start, end, service_places, ambiguous_route = _route_parts(location_text)
    start_location = location_index.resolve(start or "")
    end_location = location_index.resolve(end or "")
    route_resolved = (
        start is not None
        and end is not None
        and start_location["location_status"] == "matched"
        and end_location["location_status"] == "matched"
    )
    evidence = dict(evidence)
    evidence["is_ambiguous"] = evidence["is_ambiguous"] or ambiguous_route
    evidence["ambiguity_reason"] = (
        evidence["ambiguity_reason"]
        or ("source does not provide a unique start/end pair" if ambiguous_route else None)
    )
    evidence["requires_human_confirmation"] = not route_resolved
    dimensions = _dimensions(requirement_text)
    hard_requirements: List[Dict[str, Any]] = []
    if dimensions["min_width_m"] is not None:
        hard_requirements.append(
            {
                "requirement_type": "min_width_m",
                "value": dimensions["min_width_m"],
                "evidence_id": evidence["evidence_id"],
            }
        )
    soft_preferences: List[Dict[str, Any]] = []
    if "旅游路" in requirement_text:
        soft_preferences.append({"preference_type": "connect_to_tourism_road"})
    if "护栏" in requirement_text:
        soft_preferences.append({"preference_type": "include_guardrail"})
    return {
        "name": name,
        "source_marker": marker or None,
        "action": _action(requirement_text or name),
        "start_place": start,
        "end_place": end,
        "start_demand_id": None,
        "end_demand_id": None,
        "start_location_status": start_location["location_status"],
        "end_location_status": end_location["location_status"],
        "start_geometry": start_location["geometry"],
        "end_geometry": end_location["geometry"],
        "service_places": service_places,
        "road_class": _road_class(name),
        **dimensions,
        "priority": priority,
        "hard_requirements": hard_requirements,
        "soft_preferences": soft_preferences,
        "planning_status": "ready" if route_resolved else "blocked_needs_location_confirmation",
        "requires_human_confirmation": not route_resolved,
        "evidence": evidence,
        "supporting_evidence": [],
    }


def _workbook_projects(
    rows: Sequence[Mapping[str, Any]], location_index: ExactLocationIndex
) -> List[Dict[str, Any]]:
    projects: List[Dict[str, Any]] = []
    for row in rows:
        name = str(row["name"])
        if not ROAD_TERMS.search(_norm(name)):
            continue
        display_name = _clean(name.split("/")[-1])
        evidence = _evidence(
            Path(row["source_path"]),
            str(row["source_text"]),
            source_kind="xlsx_row",
            confidence=0.98,
            sheet=str(row["source_sheet"]),
            row=int(row["source_row"]),
            cell_range=str(row["source_cell_range"]),
        )
        locations = parse_numbered_segments(str(row["location"]))
        requirements = dict(parse_numbered_segments(str(row["note"])))
        if len(locations) == 1 and not locations[0][0]:
            requirement = f"{row['source_text']} {row['note']}"
            projects.append(
                _project_from_parts(
                    name=display_name,
                    location_text=locations[0][1],
                    requirement_text=requirement,
                    priority="high",
                    evidence=evidence,
                    location_index=location_index,
                )
            )
            continue
        for marker, location_text in locations:
            requirement = requirements.get(marker, str(row["note"]))
            projects.append(
                _project_from_parts(
                    name=display_name,
                    location_text=location_text,
                    requirement_text=requirement,
                    priority="high",
                    evidence=evidence,
                    location_index=location_index,
                    marker=marker,
                )
            )
    return projects


def _markdown_projects(
    rows: Sequence[Mapping[str, Any]], location_index: ExactLocationIndex
) -> List[Dict[str, Any]]:
    projects: List[Dict[str, Any]] = []
    for row in rows:
        name = str(row["name"])
        if not ROAD_TERMS.search(_norm(name)):
            continue
        evidence = _evidence(
            Path(row["source_path"]),
            str(row["source_text"]),
            source_kind="markdown_table_row",
            confidence=0.95,
            line=int(row["source_line"]),
        )
        locations = parse_numbered_segments(str(row["location"]))
        priority = "high" if "2022-2025" in str(row.get("timeline", "")) else "medium"
        for marker, location_text in locations:
            project = _project_from_parts(
                    name=name,
                    location_text=location_text,
                    requirement_text=f"{name} {row.get('source_text', '')}",
                    priority=priority,
                    evidence=evidence,
                    location_index=location_index,
                    marker=marker,
                )
            project["action"] = _action(name)
            projects.append(project)
    return projects


def _project_key(project: Mapping[str, Any]) -> Tuple[str, str, str]:
    return (
        _norm(project.get("road_class")),
        _norm(project.get("start_place")),
        _norm(project.get("end_place")),
    )


def _dedupe_projects(projects: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    output: List[Dict[str, Any]] = []
    keyed: Dict[Tuple[str, str, str], Dict[str, Any]] = {}
    for project in projects:
        key = _project_key(project)
        if key[1] and key[2] and key in keyed:
            keyed[key]["supporting_evidence"].append(project["evidence"])
            continue
        keyed[key] = project
        output.append(project)
    return output


def _constraint_record(
    *,
    path: Path,
    line_number: int,
    line: str,
    constraint_type: str,
    target_layer: str,
    severity: str,
    description: str,
    rule: Optional[Dict[str, Any]] = None,
    buffer_m: Optional[float] = 0.0,
    requires_confirmation: bool = False,
) -> Dict[str, Any]:
    evidence = _evidence(
        path,
        line,
        source_kind="markdown_line",
        confidence=0.97,
        line=line_number,
        requires_confirmation=requires_confirmation,
    )
    return {
        "constraint_type": constraint_type,
        "target_layer": target_layer,
        "severity": severity,
        "buffer_m": buffer_m,
        "penalty": None,
        "description": description,
        "rule": rule,
        "requires_human_confirmation": requires_confirmation,
        "evidence": evidence,
        "supporting_evidence": [],
    }


def extract_constraints_from_document(path: Path) -> List[Dict[str, Any]]:
    constraints: List[Dict[str, Any]] = []
    for line_number, raw_line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        line = _clean(re.sub(r"<[^>]+>", " ", raw_line))
        if not line:
            continue
        if "生态保护红线" in line and (STRICT_TERMS.search(line) or "不突破" in line):
            constraints.append(
                _constraint_record(
                    path=path,
                    line_number=line_number,
                    line=line,
                    constraint_type="avoid",
                    target_layer="ecological_protection",
                    severity="hard",
                    description="道路建设不得穿越生态保护红线",
                )
            )
        if "永久基本农田" in line and (STRICT_TERMS.search(line) or "不突破" in line):
            constraints.append(
                _constraint_record(
                    path=path,
                    line_number=line_number,
                    line=line,
                    constraint_type="avoid",
                    target_layer="farmland_protection",
                    severity="hard",
                    description="道路建设不得占用永久基本农田",
                )
            )
        if "地质灾害" in line and re.search(r"避免|避让", line):
            constraints.append(
                _constraint_record(
                    path=path,
                    line_number=line_number,
                    line=line,
                    constraint_type="avoid",
                    target_layer="geological_hazard_points",
                    severity="soft",
                    description="道路选线应避让地质灾害风险区",
                    buffer_m=None,
                    requires_confirmation=True,
                )
            )
        if "文物保护" in line and STRICT_TERMS.search(line):
            constraints.append(
                _constraint_record(
                    path=path,
                    line_number=line_number,
                    line=line,
                    constraint_type="avoid",
                    target_layer="historical_protection",
                    severity="hard",
                    description="道路工程不得破坏文物保护范围",
                )
            )
        if "古树名木" in line and re.search(r"避让|禁止|不得", line):
            constraints.append(
                _constraint_record(
                    path=path,
                    line_number=line_number,
                    line=line,
                    constraint_type="avoid",
                    target_layer="ancient_trees",
                    severity="hard",
                    description="道路选线应最大限度避让古树名木",
                    buffer_m=None,
                    requires_confirmation=True,
                )
            )
        if ("禁堵河" in line or "不填湖" in line) and re.search(r"河|湖|水域", line):
            constraints.append(
                _constraint_record(
                    path=path,
                    line_number=line_number,
                    line=line,
                    constraint_type="avoid",
                    target_layer="water",
                    severity="hard",
                    description="道路建设不得堵塞河道或填占湖泊水域",
                    buffer_m=0.0,
                )
            )

        road_class = None
        if "对外交通" in line or "乡道" in line:
            road_class = "township_road"
        elif "村庄主路" in line or "一级村道" in line:
            road_class = "village_main_road"
        elif "村庄支路" in line or "二级村道" in line:
            road_class = "village_branch_road"
        elif "巷道" in line:
            road_class = "lane"
        if road_class:
            width = re.search(r"(?:道路宽度(?:控制为|不小于)|宽度为|巷道不小于)\s*(\d+(?:\.\d+)?)\s*米", line)
            if width:
                value = float(width.group(1))
                constraints.append(
                    _constraint_record(
                        path=path,
                        line_number=line_number,
                        line=line,
                        constraint_type="design_standard",
                        target_layer="road_design",
                        severity="hard",
                        description=f"{road_class}最小宽度为{value:g}米",
                        rule={"metric": "min_width_m", "operator": ">=", "value": value, "road_class": road_class},
                    )
                )
            setback = re.search(r"(?:建筑物?|两侧建筑)(?:应)?后退道路?不小于\s*(\d+(?:\.\d+)?)\s*米", line)
            if setback:
                value = float(setback.group(1))
                constraints.append(
                    _constraint_record(
                        path=path,
                        line_number=line_number,
                        line=line,
                        constraint_type="design_standard",
                        target_layer="road_design",
                        severity="hard",
                        description=f"{road_class}建筑退距不小于{value:g}米",
                        rule={"metric": "building_setback_m", "operator": ">=", "value": value, "road_class": road_class},
                    )
                )
        if "消防" in line and "转弯半径" in line:
            radius = re.search(r"转弯半径不应小于\s*(\d+(?:\.\d+)?)\s*米", line)
            if radius:
                value = float(radius.group(1))
                constraints.append(
                    _constraint_record(
                        path=path,
                        line_number=line_number,
                        line=line,
                        constraint_type="design_standard",
                        target_layer="road_design",
                        severity="hard",
                        description=f"消防道路转弯半径不小于{value:g}米",
                        rule={"metric": "min_turn_radius_m", "operator": ">=", "value": value, "road_class": "fire_access"},
                    )
                )
    return constraints


def _dedupe_constraints(constraints: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    output: List[Dict[str, Any]] = []
    keyed: Dict[str, Dict[str, Any]] = {}
    for constraint in constraints:
        key = json.dumps(
            {
                "type": constraint["constraint_type"],
                "target": constraint["target_layer"],
                "severity": constraint["severity"],
                "rule": constraint["rule"],
            },
            sort_keys=True,
            ensure_ascii=False,
        )
        if key in keyed:
            keyed[key]["supporting_evidence"].append(constraint["evidence"])
            continue
        keyed[key] = constraint
        output.append(constraint)
    return output


def _dedupe_demands(demands: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    output: List[Dict[str, Any]] = []
    seen: Dict[Tuple[str, str, str], Dict[str, Any]] = {}
    for demand in demands:
        canonical_name = str(demand["name"])
        place_name = str(demand.get("place_name") or "")
        if place_name and canonical_name.endswith(f"（{place_name}）"):
            canonical_name = canonical_name[: -len(place_name) - 2]
        canonical_name = canonical_name.split("/")[-1]
        canonical_name = re.sub(r"（(?:个|米|处)）|\d+（个）", "", canonical_name)
        canonical_name = canonical_name.replace("公共厕所", "公厕").replace("公交车站点", "公交站点")
        key = (demand["demand_type"], _norm(canonical_name), _norm(place_name))
        if key in seen:
            seen[key]["supporting_evidence"].append(demand["evidence"])
            continue
        seen[key] = demand
        output.append(demand)
    return output


def _settlement_demands(
    paths: Sequence[Path], location_index: ExactLocationIndex
) -> List[Dict[str, Any]]:
    demands: List[Dict[str, Any]] = []
    seen: set[str] = set()
    for path in paths:
        for line_number, raw_line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
            line = _clean(raw_line)
            if "12个自然村" not in line or "包括" not in line:
                continue
            place_text = line.split("包括", 1)[1].split("12个自然村", 1)[0]
            place_text = place_text.replace("、和", "、").replace("和", "、").strip("，。 ：:")
            for place in _split_places(place_text):
                place = re.sub(r"自然村$", "", place)
                if not place or _norm(place) in seen:
                    continue
                seen.add(_norm(place))
                evidence = _evidence(
                    path,
                    line,
                    source_kind="markdown_line",
                    confidence=0.99,
                    line=line_number,
                )
                demands.append(
                    _make_demand(
                        name=f"{place}自然村",
                        demand_type="settlement",
                        place_name=place,
                        priority="high",
                        evidence=evidence,
                        location_index=location_index,
                        source_role="settlement_baseline",
                    )
                )
    return demands


def _assign_ids(
    demands: List[Dict[str, Any]], projects: List[Dict[str, Any]], constraints: List[Dict[str, Any]]
) -> None:
    demands.sort(key=lambda item: (item["evidence"]["source_document"], item["evidence"].get("source_row") or 0, item["name"], item.get("place_name") or ""))
    projects.sort(key=lambda item: (item["evidence"]["source_document"], item["evidence"].get("source_row") or 0, item["evidence"].get("source_line") or 0, item.get("source_marker") or "", item["name"]))
    constraints.sort(key=lambda item: (item["evidence"]["source_document"], item["evidence"].get("source_line") or 0, item["target_layer"], item["description"]))
    for index, demand in enumerate(demands, 1):
        demand["demand_id"] = f"D{index:03d}"
    for index, project in enumerate(projects, 1):
        project["project_id"] = f"RP{index:03d}"
    for index, constraint in enumerate(constraints, 1):
        constraint["constraint_id"] = f"C{index:03d}"


def _link_projects_to_demands(
    projects: Sequence[Dict[str, Any]], demands: Sequence[Dict[str, Any]]
) -> None:
    settlement_by_place: MutableMapping[str, List[str]] = defaultdict(list)
    by_place: MutableMapping[str, List[str]] = defaultdict(list)
    for demand in demands:
        if demand.get("place_name"):
            by_place[_norm(demand["place_name"])].append(demand["demand_id"])
            if demand["demand_type"] == "settlement":
                settlement_by_place[_norm(demand["place_name"])].append(demand["demand_id"])
    for project in projects:
        for end_name, field in ((project.get("start_place"), "start_demand_id"), (project.get("end_place"), "end_demand_id")):
            matches = settlement_by_place.get(_norm(end_name), []) if end_name else []
            if not matches and end_name:
                matches = by_place.get(_norm(end_name), [])
            if len(matches) == 1:
                project[field] = matches[0]


def structure_road_requirements(config: Stage2StructureConfig) -> Stage2StructureResult:
    """Run Stage 2 and write deterministic, evidence-preserving outputs."""

    config.validate()
    output_dir = config.output_dir
    output_dir.mkdir(parents=True, exist_ok=True)
    location_index = ExactLocationIndex(config.location_sources)

    workbook_rows = _extract_workbook_rows(config.demand_workbook, "demand")
    if config.current_workbook is not None:
        workbook_rows.extend(_extract_workbook_rows(config.current_workbook, "current"))
    markdown_rows = _markdown_project_rows(config.project_documents)

    demands = _dedupe_demands(
        _settlement_demands(config.constraint_documents, location_index)
        + _workbook_demands(workbook_rows, location_index)
        + _markdown_demands(markdown_rows, location_index)
    )
    projects = _dedupe_projects(
        _workbook_projects([row for row in workbook_rows if row["role"] == "demand"], location_index)
        + _markdown_projects(markdown_rows, location_index)
    )
    constraints = _dedupe_constraints(
        [
            constraint
            for path in config.constraint_documents
            for constraint in extract_constraints_from_document(path)
        ]
    )
    _assign_ids(demands, projects, constraints)
    _link_projects_to_demands(projects, demands)

    demand_features = []
    for demand in demands:
        properties = {key: value for key, value in demand.items() if key != "geometry"}
        demand_features.append(
            {"type": "Feature", "id": demand["demand_id"], "geometry": demand["geometry"], "properties": properties}
        )
    demand_payload = {
        "type": "FeatureCollection",
        "name": "demand_points",
        "crs": {"type": "name", "properties": {"name": config.display_crs}},
        "features": demand_features,
    }
    project_payload = {
        "type": "RoadProjectCollection",
        "schema_version": SCHEMA_VERSION,
        "village_name": config.village_name,
        "projects": projects,
    }
    constraint_payload = {
        "type": "RoadConstraintCollection",
        "schema_version": SCHEMA_VERSION,
        "village_name": config.village_name,
        "constraints": constraints,
    }

    review_items: List[Dict[str, Any]] = []
    unresolved_by_place: MutableMapping[str, List[Dict[str, Any]]] = defaultdict(list)
    for demand in demands:
        if demand["requires_human_confirmation"]:
            unresolved_by_place[demand.get("place_name") or demand["name"]].append(demand)
    for place_name, place_demands in sorted(unresolved_by_place.items()):
            review_items.append(
                {
                    "object_type": "demand",
                    "object_ids": [demand["demand_id"] for demand in place_demands],
                    "issue": place_demands[0]["location_reason"],
                    "question": f"请确认“{place_name}”在地图上的唯一位置。",
                }
            )
    for project in projects:
        if project["requires_human_confirmation"]:
            review_items.append(
                {
                    "object_type": "road_project",
                    "object_id": project["project_id"],
                    "issue": "起点或终点未被唯一定位，禁止进入自动规划",
                    "question": f"请在地图上确认“{project['name']}”的起点、终点及必要途经点。",
                }
            )
    for constraint in constraints:
        if constraint["requires_human_confirmation"]:
            review_items.append(
                {
                    "object_type": "constraint",
                    "object_id": constraint["constraint_id"],
                    "issue": "约束已识别，但文本没有给出可执行缓冲距离",
                    "question": f"请确认“{constraint['description']}”所需的缓冲距离或惩罚参数。",
                }
            )

    source_paths = list(
        dict.fromkeys(
            [config.demand_workbook]
            + ([config.current_workbook] if config.current_workbook else [])
            + list(config.project_documents)
            + list(config.constraint_documents)
            + list(config.location_sources)
        )
    )
    source_inventory = [
        {"path": _relative(path), "sha256": _sha256(path), "size_bytes": path.stat().st_size}
        for path in source_paths
    ]
    report = {
        "schema_version": SCHEMA_VERSION,
        "village_name": config.village_name,
        "demand_count": len(demands),
        "matched_demand_count": sum(item["location_status"] == "matched" for item in demands),
        "unresolved_demand_count": sum(item["location_status"] == "unresolved" for item in demands),
        "ambiguous_demand_count": sum(item["location_status"] == "ambiguous" for item in demands),
        "road_project_count": len(projects),
        "planning_ready_project_count": sum(item["planning_status"] == "ready" for item in projects),
        "blocked_project_count": sum(item["planning_status"] != "ready" for item in projects),
        "constraint_count": len(constraints),
        "manual_review_count": len(review_items),
        "source_inventory": source_inventory,
        "leakage_guard": {
            "planned_road_read": False,
            "planned_project_geometry_read": False,
            "location_sources": [_relative(path) for path in config.location_sources],
            "rule": "planned road/project geometry is evaluation-only and rejected as a location source",
        },
        "warnings": [
            "Unresolved or ambiguous locations have null geometry and cannot enter automatic planning.",
            "A geological-hazard/ancient-tree buffer must be confirmed before route generation.",
        ],
    }

    demand_path = output_dir / "demand_points.geojson"
    project_path = output_dir / "road_projects.json"
    constraint_path = output_dir / "road_constraints.json"
    report_path = output_dir / "stage2_report.json"
    review_path = output_dir / "manual_review_queue.json"
    inventory_path = output_dir / "source_inventory.json"
    _write_json(demand_path, demand_payload)
    _write_json(project_path, project_payload)
    _write_json(constraint_path, constraint_payload)
    _write_json(report_path, report)
    _write_json(review_path, {"type": "ManualReviewQueue", "items": review_items})
    _write_json(inventory_path, {"type": "SourceInventory", "sources": source_inventory})
    return Stage2StructureResult(
        output_dir=output_dir,
        demand_points_path=demand_path,
        road_projects_path=project_path,
        road_constraints_path=constraint_path,
        report_path=report_path,
        review_queue_path=review_path,
        report=report,
    )


__all__ = [
    "ExactLocationIndex",
    "extract_constraints_from_document",
    "parse_numbered_segments",
    "structure_road_requirements",
]
