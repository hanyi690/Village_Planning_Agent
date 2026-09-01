from pathlib import Path

from docx import Document


path = Path("F:/project/Village_Planning_Agent/docs/reports-weekly/2026-08-25-to-2026-08-29.docx")
doc = Document(path)
replacements = {
    "重构 L1/L2/L3 契约和运行时：模型输出从报告文本逐步转为可校验的事实、分析结果和规划意图。":
        "重构 L1/L2/L3 契约和运行时：把原来主要依靠 report_body 传递的信息拆成结构化阶段产物。L1 现在面向事实和观察结果，L2 面向问题、目标、需求与策略，L3 面向带约束的规划意图；同时增加了结构化输出校验、失败诊断和状态传递，便于后续程序继续处理。",
    "完善证据和空间边界：保留来源、空间范围、实体身份、置信度和诊断状态，避免证据不足时生成伪结论。":
        "完善证据和空间边界：在事实和规划对象中保留 source refs、原文片段、空间 mention、地块或设施身份、置信度以及 diagnostic/deferred 状态。对无法解析的地点、缺少候选几何或证据不足的目标，系统保持诊断或延期，不再自动猜测成可执行规划。",
    "重组 Context Runtime、Evidence Runtime、Workflow、Routing 和 Gates 等模块，明确上下文、证据、阶段编排和门禁职责。":
        "重组 Context Runtime、Evidence Runtime、Workflow、Routing 和 Gates 等模块：将模型上下文预算、证据检索、阶段调度、审批门和失败恢复分开管理，并整理 contracts、modules、infrastructure runtime 的调用边界。这样后续可以分别测试某一阶段，而不必通过整份报告判断系统是否正确。",
    "增加外部地图证据缓存、provider registry、采集验证脚本和本地 GIS 对齐审计，为多源数据接入做准备。":
        "增加外部地图证据缓存、provider registry、采集验证脚本和本地 GIS 对齐审计：目前已为 AMap、天地图等外部来源建立统一接入和缓存方向，并记录外部观察的来源与验证结果。该工作还处于接入和验证阶段，尚未等同于已经完成稳定的多源全流程规划。",
    "重新整理设计文档，区分目标设计、当前运行时、迁移台账和历史材料。":
        "重新整理设计文档：将目标架构、当前代码实际行为、迁移状态、ADR 决策和历史材料分开，新增模块级设计和认知世界模型说明。这样可以明确区分“设计上要求的能力”和“当前已经通过真实运行验证的能力”。",
}
found = set()
for paragraph in doc.paragraphs:
    text = paragraph.text.strip()
    if text in replacements:
        paragraph.clear()
        paragraph.add_run(replacements[text])
        found.add(text)
if found != set(replacements):
    missing = set(replacements) - found
    raise RuntimeError(f"missing paragraphs: {missing}")
doc.save(path)
print(f"updated={len(found)} path={path}")
