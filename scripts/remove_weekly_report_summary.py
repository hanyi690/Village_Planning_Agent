from pathlib import Path

from docx import Document


path = Path("F:/project/Village_Planning_Agent/docs/reports-weekly/2026-08-25-to-2026-08-29.docx")
doc = Document(path)
heading = "六、向老师汇报的简短表述"
paragraphs = list(doc.paragraphs)
start = next(i for i, p in enumerate(paragraphs) if p.text.strip() == heading)
end = len(paragraphs)
for i in range(start + 1, len(paragraphs)):
    if paragraphs[i].text.strip().startswith("七、"):
        end = i
        break
for paragraph in paragraphs[start:end]:
    element = paragraph._element
    element.getparent().remove(element)
for paragraph in doc.paragraphs:
    if paragraph.text.strip().startswith("七、依据与阅读入口"):
        for run in paragraph.runs:
            run.text = run.text.replace("七、依据与阅读入口", "六、依据与阅读入口")
doc.save(path)
print(f"removed_paragraphs={end-start} path={path}")
