from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "reports-weekly" / "2026-08-25-to-2026-08-29.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.5)
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.6)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(5)
for name, size, color in [("Heading 1", 14, "1F4E79"), ("Heading 2", 11.5, "2F75B5")]:
    style = styles[name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(4)
run = title.add_run("项目周报：规划分析链路结构化升级")
run.bold = True
run.font.size = Pt(20)
run.font.color.rgb = RGBColor(31, 78, 121)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(14)
sr = subtitle.add_run("2026 年 8 月 25 日—8 月 29 日｜Village Planning Agent")
sr.font.size = Pt(10)
sr.font.color.rgb = RGBColor(100, 100, 100)

lead = doc.add_paragraph()
lead.paragraph_format.space_after = Pt(10)
lr = lead.add_run("本周结论：")
lr.bold = True
lr.font.color.rgb = RGBColor(31, 78, 121)
lead.add_run(
    "项目从“资料直接生成报告”的模式，进一步转向“资料理解—证据组织—分层分析—空间接地—方案交付”的结构化规划系统。"
    "L1、L2、L3 已分别明确为事实构建、规划认知和规划意图三个阶段。"
)

add_heading(doc, "一、与前期工作的衔接")
doc.add_paragraph(
    "前期已经验证过单地块及同源地块窄切片的空间规划闭环：从 ScenarioSet、执行审批，到 EXECUTE、VALIDATE、RECONCILE 和 DELIVER，"
    "系统能够将明确的局部目标落到空间候选、执行地块和交付图层。"
)
doc.add_paragraph(
    "本周工作的重点前移到这条执行链路的上游：让资料、证据和规划分析形成可复用的结构化产物。需要说明的是，前期闭环主要是局部窄切片，"
    "并不等同于全村法定规划完成；本周的架构重构也仍需通过真实全流程运行进一步验证。"
)

add_heading(doc, "二、本周主要工作")
add_bullets(doc, [
    "重构 L1/L2/L3 契约和运行时：模型输出从报告文本逐步转为可校验的事实、分析结果和规划意图。",
    "完善证据和空间边界：保留来源、空间范围、实体身份、置信度和诊断状态，避免证据不足时生成伪结论。",
    "重组 Context Runtime、Evidence Runtime、Workflow、Routing 和 Gates 等模块，明确上下文、证据、阶段编排和门禁职责。",
    "增加外部地图证据缓存、provider registry、采集验证脚本和本地 GIS 对齐审计，为多源数据接入做准备。",
    "重新整理设计文档，区分目标设计、当前运行时、迁移台账和历史材料。",
])

add_heading(doc, "三、系统模块简述")
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
table.autofit = False
widths = [Cm(3.4), Cm(6.2), Cm(7.0)]
for cell, text in zip(table.rows[0].cells, ["模块", "主要职责", "当前阶段"]):
    set_cell_text(cell, text, bold=True, color=(255, 255, 255))
    shade(cell, "1F4E79")
for row in [
    ("DATA / Source Cognition", "接收报告、表格、GIS 和外部服务，识别字段、来源和空间范围。", "外部证据缓存和 provider 接口已开始形成。"),
    ("Evidence / KG", "组织证据、实体、地点、关系和来源 lineage。", "检索、图谱和证据边界持续收敛。"),
    ("L1", "将资料和证据整理为带类型、来源、范围和状态的结构化事实。", "契约和观察链路已基本建立。"),
    ("L2", "基于事实形成问题、目标、需求、适宜性和策略。", "正在从文本分析转为结构化规划认知。"),
    ("L3", "把规划认知编译为带约束、空间要求和 lineage 的规划意图。", "正在完善 typed claim 和 Planning Program。"),
    ("GROUND / SELECT / EXECUTE", "将规划意图落到候选空间，选择方案并形成执行对象。", "窄切片已有闭环，复杂场景仍需验证。"),
    ("DELIVER", "输出报告、地图、规划图层和诊断信息。", "已有局部交付，尚未稳定形成全村规划图。"),
]:
    cells = table.add_row().cells
    for cell, text in zip(cells, row):
        set_cell_text(cell, text)
for row in table.rows:
    for cell, width in zip(row.cells, widths):
        cell.width = width

add_heading(doc, "四、当前产出与不足")
doc.add_paragraph(
    "本周已经完成了结构化架构和模块边界的一轮收敛，测试和代码层面能够验证若干局部能力。"
    "但是，老师最关心的最终规划产出图，目前还没有在新架构下稳定跑通全流程。当前更准确的判断是："
)
add_bullets(doc, [
    "已有：明确局部目标的窄切片可以形成空间候选、执行结果和规划图层。",
    "已有：L1/L2/L3、证据和空间对象之间的契约、来源和诊断边界正在建立。",
    "不足：多目标、多设施、道路网络、覆盖和拓扑分区尚未形成稳定的全流程规划图。",
    "风险：结构化模块已经较多，但如果没有一张新的、可重复生成的规划产出图，系统价值仍不够直观。",
])

note = doc.add_paragraph()
note.paragraph_format.left_indent = Cm(0.4)
note.paragraph_format.right_indent = Cm(0.4)
note.paragraph_format.space_before = Pt(4)
note.paragraph_format.space_after = Pt(8)
nr = note.add_run("阶段判断：")
nr.bold = True
nr.font.color.rgb = RGBColor(192, 80, 77)
note.add_run("当前不是“全流程已经完成”，而是“结构化分析系统已基本搭起，下一阶段必须用一个可控样例重新打通并产出可信规划图”。")

add_heading(doc, "五、下一步计划")
for text in [
    "选定一个目标明确、空间范围可控的村庄样例，固定输入资料、GIS 图层和规划目标。",
    "先确保 L1 事实、L2 分析、L3 规划意图都能生成并保留来源，再进入空间接地。",
    "优先打通一个可展示的 site-area 或设施选址切片，生成规划 footprint、余量和地图图层。",
    "用事件、结构化产物和截图同步验收，不再只用最终报告文本判断成功。",
    "在局部闭环稳定后，再扩展到多目标、道路、覆盖和拓扑分区，并形成新的全流程规划图。",
]:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)

add_heading(doc, "六、向老师汇报的简短表述")
doc.add_paragraph(
    "前期我们已经把一个明确的局部规划目标跑到了空间交付，本周的重点是把上游的资料理解、证据组织以及 L1、L2、L3 的规划分析进一步结构化。"
    "现在 L1 负责形成有来源的事实，L2 负责问题、目标和策略分析，L3 负责形成带约束的规划意图，之后再接到空间接地、方案选择和执行。"
    "目前新架构已经基本搭建，但全流程规划图还没有稳定产出。下一步会固定一个可控样例，优先重新打通一条可展示的端到端链路，"
    "先拿到可信的局部规划图，再逐步扩展到多目标、道路、覆盖和拓扑规划。"
)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Village Planning Agent｜阶段性周报").font.size = Pt(8)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
