from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


path = Path("F:/project/Village_Planning_Agent/docs/reports-weekly/2026-08-25-to-2026-08-29.docx")
doc = Document(path)
for paragraph in doc.paragraphs:
    if paragraph.text.startswith("阶段判断：阶段判断："):
        paragraph.clear()
        paragraph.add_run("阶段判断：目前已完成结构化分析系统的主要架构整理，下一阶段将继续通过可控样例完善端到端验证，并形成更稳定的规划图输出。")
        break
heading = doc.add_heading("六、依据与阅读入口", level=1)
heading.paragraph_format.keep_with_next = True
for label, target in [
    ("目标系统总设计", "docs/design/00_SYSTEM_DESIGN.md"),
    ("L1 契约与不变量", "docs/design/01_CONTRACTS_AND_INVARIANTS.md"),
    ("L2 规划认知", "docs/design/02_PLANNING_COGNITION.md"),
    ("L3 规划意图", "docs/design/03_PLANNING_INTENT.md"),
    ("证据系统", "docs/design/04_EVIDENCE.md"),
    ("投影与空间接地", "docs/design/05_PROJECTION_AND_GROUND.md"),
    ("当前运行时", "docs/architecture/CURRENT_RUNTIME.md"),
    ("迁移台账", "docs/design/05_MIGRATION.md"),
]:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(f"{label}：{target}")
doc.save(path)
print(f"repaired={path}")
