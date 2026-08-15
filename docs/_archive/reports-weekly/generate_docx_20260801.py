"""Generate weekly report docx for 2026-08-01 (with rendered map)."""
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os

BASE = r"F:\project\Village_Planning_Agent"
MAP_IMG = os.path.join(BASE, r"docs\gis\output\integrated_planning_map.png")
OUTPUT = os.path.join(BASE, r"docs\weekly-reports\2026-08-01.docx")
os.makedirs(os.path.dirname(OUTPUT), exist_ok=True)

doc = Document()

style = doc.styles["Normal"]
font = style.font
font.name = "宋体"
font.size = Pt(12)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
pf = style.paragraph_format
pf.space_after = Pt(4)
pf.line_spacing = Pt(20)


def h1(text):
    t = doc.add_heading(text, level=1)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in t.runs:
        run.font.size = Pt(16)
        run.font.name = "黑体"
        run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def h2(text):
    t = doc.add_heading(text, level=2)
    for run in t.runs:
        run.font.size = Pt(13)
        run.font.name = "黑体"
        run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def para(text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
    p.add_run(text)
    return p


# ── 标题 ──
title = doc.add_heading("周报（2026-07-28 ~ 2026-08-01）", level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.size = Pt(18)
    run.font.name = "黑体"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

# ── 一、实现了什么 ──
h2("一、实现了什么")
para(
    "构建 EXPLORE 统一访问层（6 工具 + 空间认知概要注入维度分析 prompt），"
    "_optimize_semantics 升级为 L1/L2/L3 均查 KG，修复 RAG 3 个真实 bug，451 项测试通过。",
    bold_prefix="1. EXPLORE 统一访问层：",
)
para(
    "BridgeAdvisor 空间认知 brief 注入 L1 prompt（村内/村外铁律：4 个村外行政村注记点实测在村界外，绝不绑定）；"
    "wave 末尾 PARSE 提取实体语义注解回流 KG（村名→居民点聚类→BSM，三重 write-gate）；"
    "语义记忆轮次持久化（跨轮 diff/supersede）。",
    bold_prefix="2. 报告→KG 语义闭环：",
)
para(
    "L3 报告分区表 → 动态词汇表（LLM 提取 + 表格闭集防编造）→ PARSE 表格解析 → Binder 语义绑定 → "
    "EXECUTE 生成 L4 FunctionalZone（多因子规则 + 三线拦截 + 面积对齐）→ PNG 出图。",
    bold_prefix="3. 产业空间布局规划图（自定义功能区）：",
)
para(
    "修复 aiosqlite≥0.22 进程挂死（checkpointer 裸连接 + engine 引用计数陷阱），"
    "shutdown_database + watchdog 强杀兜底，崩溃/正常路径均自行退出。",
    bold_prefix="4. 进程自结束机制：",
)
para(
    "upload_flow 47 例（修 2 个 DB bug）、report_store 3 bug、Pipeline E2E 4 阶段；"
    "全量 520 passed，35 个失败全部 pre-existing（零新增）。",
    bold_prefix="5. 测试与硬化：",
)

# ── 二、实现思路 ──
h2("二、实现思路")
para(
    "村名→居民点聚类→候选 BSM 的语义绑定；三线/保护地块拦截由 KG 上下文驱动；"
    "L4 功能区带 provenance（derived_from/typed_by）可回溯。",
    bold_prefix="KG 是唯一空间语义权威：",
)
para(
    "功能区类别随报告变化（LLM 自由给色），但表格闭集（LLM 只增强报告表格已有的分区，编造类别丢弃），"
    "既动态又防幻觉。",
    bold_prefix="词汇表 = L3 报告动态产物：",
)
para(
    "LLM 只产类别定义与规则，不产坐标；地块选择走多因子规则 DSL（地类/三线/坡度/道路距离/POI 邻近 + and/any/not），"
    "面积按报告目标截断对齐。",
    bold_prefix="几何 100% 确定性：",
)
para(
    "连接显式关闭（治本）→ watchdog 强杀（兜底）→ 环节超时（防卡）→ 状态断点（可续跑）。",
    bold_prefix="进程自结束四层防御：",
)

# ── 三、实现成果（图） ──
h2("三、实现成果")
p = doc.add_paragraph("产业空间布局规划图（660 地块现状底图 + 4 大功能分区 + 三线 + 道路 + 17 个项目点 + 图例）")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
pic = doc.add_paragraph()
pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
if os.path.exists(MAP_IMG):
    pic.add_run().add_picture(MAP_IMG, width=Inches(6.3))

para(
    "基本农田保护区（155 地块，耕地 1100，面积 78.83 ha 截断对齐）；一般农业区（88 地块）；"
    "村庄建设区（111 地块，宅基地 2121/2122，图斑级剔除三线重叠）；"
    "生态保护区（65 地块，林地 1340/1350，is_protection 跳过拦截）。",
    bold_prefix="自定义功能区层（L3 报告真实分区，零编造）：",
)
para(
    "docs/gis/output/integrated_planning_map.png、custom_functional_zone.geojson（4 区 + 全图层 GeoJSON）",
    bold_prefix="产出文件：",
)

# ── 四、未来优化 ──
h2("四、未来优化")
para(
    "坡度（DEM）、道路距离、POI 邻近因子注入地块上下文；LLM 提取复合规则表达式。",
    bold_prefix="1. 多因子完整接入：",
)
para("兼容矩阵 + 设施落位器 + validator 过红线。", bold_prefix="2. 项目点落位：")
para("补建设用地总量/服务半径约束，validation 3 轮收敛。", bold_prefix="3. MIP 校验收敛：")
para("图例按用地分类指南配色、CGCS2000 声明、图名/制图时间落款。", bold_prefix="4. 渲染规范对齐：")

doc.save(OUTPUT)
print(f"Saved: {OUTPUT}")
