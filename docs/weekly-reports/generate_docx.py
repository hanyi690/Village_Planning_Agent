"""Generate weekly report docx for 2026-07-25."""
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os

BASE = r"F:\project\Village_Planning_Agent"
IMG_DIR = os.path.join(BASE, r"docs\gis\kg_visualization")
OUTPUT = os.path.join(BASE, r"docs\weekly-reports\2026-07-25.docx")

doc = Document()

# ── 全局默认字体 ──
style = doc.styles["Normal"]
font = style.font
font.name = "宋体"
font.size = Pt(14)  # 四号
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
pf = style.paragraph_format
pf.space_after = Pt(4)
pf.line_spacing = Pt(22)

# ── 标题 ──
title = doc.add_heading("周报（2026-07-21 ~ 2026-07-25）", level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.size = Pt(18)
    run.font.name = "黑体"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

# ── 1. KG 四层架构 ──
doc.add_heading("1. 知识图谱四层架构 + 增量设计（v4.4）", level=2)
p = doc.add_paragraph()
p.add_run(
    "知识图谱从单层结构重构为四层语义架构：L1 现状层（660 个地块、175 条道路、44 个设施），"
    "L2 约束层（耕地/生态/历史保护 3 条管控规则），L3 意图层（路网连通推理），"
    "L4 规划变更层（PlanBlock 增量设计）。"
)

p = doc.add_paragraph()
p.add_run("关键设计思路：").bold = True
p.add_run(
    "原有方案 PlannedState 为每轮规划生成完整镜像，数据膨胀 31%。"
    "改用 PlanBlock 增量模式，只存变更 delta（37 个三元组 vs 原方案约 6,600 个），"
    "支持 split（拆分）、merge（合并）、reclassify（重分类）三种基本操作，"
    "通过 derived_from 链接追溯来源。"
)

p = doc.add_paragraph()
p.add_run("验证：").bold = True
p.add_run(
    "金田村实际数据构建 21,253 个三元组，73 项自动化测试全部通过，"
    "含 SPARQL 查询验证、路网连通查询、L1→L4 有效用途回退。"
)

# 插入统计图
img_stat = os.path.join(IMG_DIR, "kg_statistics.png")
if os.path.exists(img_stat):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(img_stat, width=Inches(5.2))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("▲ 左：实体分布柱状图；右：三层空间关系分布").font.size = Pt(10.5)

# ── 2. Pipeline 跑通 ──
doc.add_heading("2. 完整管线跑通 + 性能优化", level=2)
p = doc.add_paragraph()
p.add_run(
    "金田村完整管线（加载数据→构建 KG→MIP 分配→生成图层→渲染出图）"
    "经两轮实际运行验证。7-24 相比 7-23 首次运行：总耗时从 63 秒降至 35.6 秒（-43%），"
    "MIP 分配从 5.4 秒降至 2.7 秒（-50%），渲染从 11.1 秒降至 2.9 秒（-74%）。"
    "约束从 1 条扩展至 4 条，验证检查从 8 项增至 9 项（含 VNE 合规检查）。"
    "最终产出 11 个 GeoJSON 图层和 3 张渲染图，设施点从注记图层中独立拆分。"
)

# ── 3. Schema 抽象 ──
doc.add_heading("3. Schema 层数据模型重构", level=2)
p = doc.add_paragraph()
p.add_run(
    '数据模型从“以图层语义为中心”重构为“以数据结构为中心”：'
    'SpatialFeature 下新增 PointFeature / LineFeature / PolygonFeature 三个中间基类，'
    '构造时强制校验几何类型；地块拆为 current_landuse 和 planned_landuse 两个独立字段；'
    'GISDataBundle 升级为 ProjectDataView（多村）+ VillageBundle（单村）两层结构。'
    '对照 GB50445、GB/T 21010、QMDLTB 等 21 份规范，补充了 10 个缺失字段和 6 类设施子类型。'
)

# ── 4. Data 双轨 ──
doc.add_heading("4. Data 层双轨设计", level=2)
p = doc.add_paragraph()
p.add_run(
    "已注册类型走严格校验（28 个 layer_type 映射），未注册类型自动推断"
    "（几何类型分桶→字段名知识库匹配→编码值查表→动态模型生成）。"
    "栅格和表格各有独立推断流程。配置拆分为 sources.yaml（物理连接）"
    "和 projects.yaml（项目逻辑），避免混改。"
)

# ── 下周计划 ──
doc.add_heading("下周计划", level=2)
p = doc.add_paragraph()
p.add_run(
    "进入代码实现阶段：先完成 Schema 类型层编码（S1-S10），"
    "再做 v4.3 规范补全（S11-S15），然后进行 VillageBundle/ProjectDataView 重组"
    "及各下游模块的代码大修。"
)

doc.save(OUTPUT)
print(f"Saved to {OUTPUT}")